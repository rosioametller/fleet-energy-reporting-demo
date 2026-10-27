from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from .domain import parse_day

def safe_cell(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'"+value
    return value

def decorate(sheet, table_name):
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    for c in sheet[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='234A42')
        c.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 28
    for col in sheet.columns:
        letter = col[0].column_letter
        sheet.column_dimensions[letter].width = min(65, max(18, max(len(str(c.value or '')) for c in col)+3))
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True)
        sheet.add_table(table)

def build_workbook(snapshot, rows, warnings, advanced=False):
    book = Workbook(); report = book.active; report.title = 'Reporte'
    report.append(['Vehículo ficticio', 'Identificador interno', 'Fecha', 'Distancia (km)', 'Energía (kWh)', 'Actividad'])
    units = {u['id']: u for u in snapshot['units']}
    for r in rows:
        u = units[r['unit_id']]
        report.append([safe_cell(u['plate']), safe_cell(u['internal_key']), r['day'], r['distance'], r['energy'],
            'Cero confirmado' if r['zero_confirmed'] else 'Actividad registrada'])
    for row in report.iter_rows(min_row=2):
        row[2].number_format = 'yyyy-mm-dd'
        row[3].number_format = row[4].number_format = '0.00'
    decorate(report, 'LabReport')
    summary = book.create_sheet('Resumen')
    summary.append(['Indicador', 'Valor'])
    summary.append(['Reporte', safe_cell(snapshot['series_name'])]); summary.append(['Origen', 'Mediciones generadas desde cero para esta DEMO'])
    summary.append(['Desde', parse_day(snapshot['start'])]); summary.append(['Hasta', parse_day(snapshot['end'])])
    summary.append(['Vehículos', len(units)]); summary.append(['Observaciones', len(rows)])
    summary.append(['Distancia total (km)', sum(r['distance'] for r in rows)])
    summary.append(['Energía total (kWh)', sum(r['energy'] for r in rows)])
    for row in summary.iter_rows(min_row=2):
        if row[0].value in {'Desde','Hasta'}: row[1].number_format = 'yyyy-mm-dd'
    decorate(summary, 'LabSummary')
    if warnings:
        notes = book.create_sheet('Advertencias'); notes.append(['Observación'])
        for message in warnings: notes.append([safe_cell(message)])
        decorate(notes, 'LabWarnings')
    if advanced:
        technical = book.create_sheet('Detalle técnico'); technical.append(['Dato', 'Valor'])
        technical.append(['Escenario', snapshot['scenario']]); technical.append(['Formato', 'DEMO sintética'])
        decorate(technical, 'LabTechnical')
    book.properties.creator = 'Fleet Energy Lab'; book.properties.title = 'Reporte sintético de flota'
    output = BytesIO(); book.save(output); book.close(); return output.getvalue()
