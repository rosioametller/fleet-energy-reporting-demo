import argparse
from datetime import datetime
import hashlib
import html
import json
from pathlib import Path
import sys
import tempfile
import socket
import threading
from urllib.parse import urlsplit
from io import BytesIO

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.dont_write_bytecode=True
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook
from fleet_lab.web import create_lab_server
from fleet_lab.repository import LabDatabase
from fleet_lab.application import FleetWorkspace


def run_review(output, pictures):
    output.mkdir(parents=True,exist_ok=True);pictures.mkdir(parents=True,exist_ok=True)
    steps=[];errors=[];external=[]
    with tempfile.TemporaryDirectory(prefix='fleet-browser-') as folder, socket.socket() as offline_gateway:
        offline_gateway.bind(('127.0.0.1',0))
        # Bound but not listening: no usable HTTP(S) proxy, except bypassed loopback.
        proxy_address=f'http://127.0.0.1:{offline_gateway.getsockname()[1]}'
        server=create_lab_server(folder);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        origin=f'http://127.0.0.1:{server.server_port}'
        try:
            with sync_playwright() as pw:
                browser=pw.chromium.launch(headless=True,args=['--disable-background-networking','--disable-component-update'],proxy={'server':proxy_address,'bypass':'127.0.0.1,localhost'})
                try:
                    context=browser.new_context(accept_downloads=True,viewport={'width':1440,'height':1000})
                    def gate(route):
                        target=urlsplit(route.request.url)
                        if target.scheme=='http' and target.netloc==urlsplit(origin).netloc:route.continue_()
                        else:external.append('blocked external request');route.abort()
                    context.route('**/*',gate)
                    page=context.new_page();page.on('pageerror',lambda e:errors.append(str(e)))
                    def navigate(panel):
                        page.locator(f'[data-panel="{panel}"]').click()
                        expect(page.locator(f'[data-page="{panel}"]')).to_be_visible()
                    def photo(name):page.screenshot(path=str(pictures/(name+'.png')),full_page=True)
                    page.goto(origin)
                    expect(page.get_by_role('heading',name='De una flota ficticia a un reporte útil.')).to_be_visible()
                    expect(page.locator('#lab-stats')).to_contain_text('6');photo('inicio')
                    steps.append('Inicio: seis unidades sintéticas, tres reportes y estado local')
                    navigate('catalog');expect(page.locator('#lab-catalog-list')).to_contain_text('DEMO-901');photo('vehiculos')
                    for key,value in {'lab-plate':'DEMO-950','lab-internal':'LAB-7950','lab-lookup':'source-unit-90','lab-label':'Unidad experimental nueva'}.items():page.locator('#'+key).fill(value)
                    page.get_by_role('button',name='Guardar unidad',exact=True).click();expect(page.locator('#lab-message')).to_contain_text('Unidad ficticia guardada')
                    page.locator('#lab-search').fill('source-unit-90');expect(page.locator('#lab-catalog-list tbody tr')).to_have_count(1)
                    row=page.locator('#lab-catalog-list tbody tr')
                    page.once('dialog',lambda dialog:dialog.accept());row.get_by_role('button',name='Desactivar',exact=True).click();expect(row.locator('.status-tag')).to_have_text('Inactiva')
                    page.once('dialog',lambda dialog:dialog.accept());row.get_by_role('button',name='Reactivar',exact=True).click();expect(row.locator('.status-tag')).to_have_text('Activa')
                    steps.append('Alta, búsqueda, desactivación confirmada y reactivación con identidades separadas')
                    navigate('compose');expect(page.locator('#lab-eligible')).not_to_contain_text('DEMO-950')
                    steps.append('Alta de catálogo sin asociación automática')
                    navigate('catalog')
                    new=next(u for u in server.workspace.db.catalog()['units'] if u['plate']=='DEMO-950')
                    page.locator('#lab-link-unit').select_option(new['id']);page.locator('#lab-link-series').select_option('boreal')
                    page.get_by_role('button',name='Guardar asociación',exact=True).click();expect(page.locator('#lab-message')).to_contain_text('Asociación y vigencia guardadas')
                    page.locator('#lab-search').fill('');photo('grupos')
                    steps.append('Asociación explícita a grupo/reporte con vigencia')
                    navigate('compose');page.locator('#lab-series').select_option('boreal');expect(page.locator('#lab-eligible')).to_contain_text('DEMO-950')
                    for checkbox in page.locator('input[name="lab-unit"]').all():checkbox.uncheck()
                    page.locator('input[name="lab-unit"][value="'+new['id']+'"]').check()
                    assert not page.locator('#lab-advanced').is_checked()
                    assert page.locator('#lab-scenario').input_value()=='steady'
                    page.locator('#lab-scenario').select_option('mixed')
                    page.get_by_role('button',name='Revisar vista previa',exact=True).click();expect(page.locator('#lab-preview')).to_contain_text('Período incompleto')
                    expect(page.locator('#lab-preview')).to_contain_text('Cero confirmado')
                    page.locator('#lab-scenario').select_option('steady')
                    assert page.locator('#lab-generate-actions').is_hidden()
                    expect(page.locator('#lab-preview')).to_contain_text('Revisa la selección actual')
                    page.locator('#lab-scenario').select_option('mixed')
                    page.get_by_role('button',name='Revisar vista previa',exact=True).click()
                    expect(page.locator('#lab-preview')).to_contain_text('Período incompleto')
                    steps.append('Cambiar el alcance descarta la vista previa anterior y exige recalcular')
                    photo('crear-reporte')
                    page.get_by_role('button',name='Generar Excel',exact=True).click();expect(page.locator('#lab-generated')).to_be_visible()
                    with page.expect_download() as future:page.locator('#lab-download').click()
                    download=future.value;download.save_as(output/'reporte_sintetico.xlsx')
                    content=(output/'reporte_sintetico.xlsx').read_bytes();book=load_workbook(BytesIO(content))
                    try:
                        assert book.sheetnames==['Reporte','Resumen','Advertencias']
                        assert book['Reporte'].max_row==7
                        assert book['Reporte']['C2'].data_type=='d' and book['Reporte']['D2'].data_type=='n'
                        assert any(book['Reporte'].cell(r,4).value==0 for r in range(2,8))
                        assert len(book['Reporte'].tables)==1
                        sheet_names=book.sheetnames
                        # Visualize actual workbook cells in an isolated, blank page.
                        panels=[]
                        for sheet in book:
                            lines=['<table>']
                            for i,record in enumerate(sheet.values):
                                tag='th' if i==0 else 'td'
                                lines.append('<tr>'+''.join('<'+tag+'>'+html.escape(v.strftime('%Y-%m-%d') if isinstance(v,datetime) else str(v if v is not None else ''))+'</'+tag+'>' for v in record)+'</tr>')
                            lines.append('</table>');panels.append('<h2>'+html.escape(sheet.title)+'</h2>'+''.join(lines))
                    finally:book.close()
                    steps.append('Generación, descarga real y reapertura: Reporte, Resumen y Advertencias')
                    navigate('history');expect(page.locator('#lab-history')).to_contain_text('DEMO-950')
                    page.get_by_role('button',name='Aceptar DEMO',exact=True).click();expect(page.locator('#lab-history')).to_contain_text('Aceptado en DEMO');photo('historial')
                    saved=server.workspace.timeline()[0];assert saved['sha256']==hashlib.sha256(content).hexdigest();assert saved['decision']=='accepted'
                    steps.append('Decisión DEMO asociada al SHA-256 exacto de la descarga')
                    navigate('catalog');page.locator('#lab-backup').click();expect(page.locator('#lab-backup-info')).to_contain_text('SHA-256')
                    before=server.workspace.db.catalog()
                    with page.expect_download() as future:page.get_by_role('link',name='Exportar catálogo JSON',exact=True).click()
                    future.value.save_as(output/'catalogo_sintetico.json')
                    imported=json.loads((output/'catalogo_sintetico.json').read_bytes());imported['units'][0]['label']='Unidad importada ficticia'
                    page.locator('#lab-import-file').set_input_files({'name':'catalogo_demo.json','mimeType':'application/json','buffer':json.dumps(imported).encode()})
                    page.locator('#lab-import-preview').click();expect(page.locator('#lab-import-result')).to_contain_text('No se ha aplicado ningún cambio')
                    assert server.workspace.db.catalog()==before
                    page.once('dialog',lambda dialog:dialog.accept());page.locator('#lab-import-apply').click();expect(page.locator('#lab-message')).to_contain_text('Catálogo importado')
                    assert server.workspace.db.catalog()['units'][0]['label']=='Unidad importada ficticia'
                    steps.append('Exportación, vista previa de importación sin escritura y confirmación de importación')
                    page.once('dialog',lambda dialog:dialog.accept());page.locator('#lab-restore').click();expect(page.locator('#lab-message')).to_contain_text('Catálogo restaurado')
                    assert server.workspace.db.catalog()==before and len(server.workspace.timeline())==1
                    assert any(e['kind']=='catalog_restored' for e in server.workspace.db.events())
                    steps.append('Respaldo con hash y restauración sin borrar historial ni auditoría')
                    navigate('system');expect(page.locator('#lab-system')).to_contain_text('Proveedor local simulado');expect(page.locator('#lab-audit')).to_contain_text('catalog_restored');photo('sistema')
                    steps.append('Estado del sistema y auditoría sin rutas personales')
                    navigate('compose');page.locator('#lab-scenario').select_option('conflict')
                    page.get_by_role('button',name='Revisar vista previa',exact=True).click();expect(page.locator('#lab-error')).to_contain_text('Conflicto')
                    assert page.locator('#lab-generate-actions').is_hidden();assert len(server.workspace.timeline())==1
                    steps.append('Conflicto bloquea la generación sin crear un archivo ni un historial falso')
                    for panel in ['home','catalog','compose','history','system']:
                        page.set_viewport_size({'width':390,'height':844});navigate(panel)
                        assert page.evaluate('document.documentElement.scrollWidth<=window.innerWidth+1'),panel
                    steps.append('Cinco pantallas utilizables a 390 px sin desbordamiento horizontal de página')
                    visual=context.new_page();visual.set_viewport_size({'width':1440,'height':1100})
                    visual.set_content('<!doctype html><html lang="es"><meta charset="utf-8"><title>Excel sintético</title><style>body{font:14px Segoe UI,sans-serif;padding:34px;background:#f3f2e9;color:#18372f}table{border-collapse:collapse;background:white;width:100%;margin-bottom:26px}td,th{border:1px solid #d9dfd4;padding:10px;text-align:left}th{background:#234a42;color:white}h1{font-family:Georgia,serif;font-weight:400}h2{font-size:18px}</style><h1>Excel simplificado · Datos inventados</h1><p>Vista reconstruida desde el archivo XLSX descargado.</p>'+''.join(panels)+'</html>')
                    visual.screenshot(path=str(pictures/'excel-simplificado.png'),full_page=True)
                    assert not errors and not external
                    context.close()
                finally:browser.close()
        finally:server.shutdown();server.server_close();thread.join(5)
        assert not thread.is_alive()
        reopened=FleetWorkspace(LabDatabase(folder));assert len(reopened.timeline())==1 and len(reopened.db.catalog()['units'])==7
        steps.append('Reinicio conserva catálogo e historial; servidor y navegador propios cerrados')
        base_request=dict(series_id='boreal',unit_ids=['veh-901'],start='2032-04-01',end='2032-04-07',scenario='steady',advanced=False)
        baseline=reopened.produce_run(base_request)
        basic=reopened.obtain_file(baseline['id']);(output/'reporte_basico.xlsx').write_bytes(basic)
        baseline_book=load_workbook(BytesIO(basic))
        try:assert baseline_book.sheetnames==['Reporte','Resumen']
        finally:baseline_book.close()
        steps.append('Salida predeterminada sin advertencias: solo Reporte y Resumen')

    result=dict(flows=len(steps),steps=steps,javascript_errors=errors,external_requests=external,downloads=2,
        offline_gateway_guard=True,excel_files=2,default_excel_sheets=['Reporte','Resumen'],excel_sheets=sheet_names,excel_sha256=hashlib.sha256(content).hexdigest(),server_closed=True)
    (output/'browser-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(flows=len(steps),errors=len(errors),external=len(external),excel_files=2)))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=ROOT/'.qa/review');parser.add_argument('--screenshots',type=Path,default=ROOT/'screenshots');args=parser.parse_args()
    run_review(args.output,args.screenshots)
