from copy import deepcopy
from datetime import date
from io import BytesIO
import hashlib
import http.client
import json
from pathlib import Path
import sqlite3
import socket
from unittest.mock import patch, MagicMock
import tempfile
import threading
import unittest
from openpyxl import load_workbook
from fleet_lab.application import FleetWorkspace
from fleet_lab.domain import InputFault, inclusive_days, shape_unit, normalize_observations
from fleet_lab.repository import LabDatabase
from fleet_lab.providers import LaboratorySource
from fleet_lab.web import create_lab_server

class IsolatedLab(unittest.TestCase):
    def setUp(self):
        self.sandbox=tempfile.TemporaryDirectory(prefix='fleet-unit-');self.addCleanup(self.sandbox.cleanup)
        self.folder=Path(self.sandbox.name);self.db=LabDatabase(self.folder);self.engine=FleetWorkspace(self.db)
        self.payload=dict(series_id='boreal',unit_ids=['veh-901'],start='2032-04-01',end='2032-04-07',scenario='steady',advanced=False)

class Rules(IsolatedLab):
    def test_seed_is_new_and_synthetic(self):
        items=self.db.catalog()['units'];self.assertEqual(len(items),6)
        self.assertTrue(all(u['plate'].startswith('DEMO-9') for u in items))
        self.assertTrue(all(u['internal_key']!=u['plate'] and u['lookup_key']!=u['internal_key'] for u in items))
    def test_invalid_plate_cannot_enter_catalog(self):
        unit=deepcopy(self.db.catalog()['units'][0]);unit['plate']='NOT-A-DEMO-PLATE'
        with self.assertRaises(InputFault):shape_unit(unit)
    def test_large_or_reversed_period_rejected(self):
        for a,b in [('2032-04-07','2032-04-01'),('2032-01-01','2032-07-01'),('bad','2032-04-01')]:
            with self.subTest(a=a),self.assertRaises(InputFault):inclusive_days(a,b)
    def test_zero_requires_confirmation(self):
        units=self.db.eligible('boreal','2032-04-01','2032-04-01');day=date(2032,4,1)
        with self.assertRaises(InputFault):normalize_observations([dict(unit_id=units[0]['id'],day=day,distance=0,energy=0,zero_confirmed=False)],units,[day])
    def test_non_finite_negative_and_boolean_metrics_rejected(self):
        units=[self.db.eligible('boreal','2032-04-01','2032-04-01')[0]];day=date(2032,4,1)
        for value in (-1,float('nan'),float('inf'),True):
            with self.subTest(value=value),self.assertRaises(InputFault):normalize_observations([dict(unit_id=units[0]['id'],day=day,distance=value,energy=1,zero_confirmed=False)],units,[day])
    def test_confirmed_zero_is_preserved(self):
        _,rows,_=self.engine.stage_run(dict(self.payload,scenario='zero'))
        zero=[r for r in rows if r['zero_confirmed']];self.assertEqual(len(zero),1);self.assertEqual(zero[0]['distance'],0)
    def test_missing_is_not_replaced_with_zero(self):
        _,rows,notes=self.engine.stage_run(dict(self.payload,scenario='missing'))
        self.assertEqual(len(rows),6);self.assertTrue(notes);self.assertNotIn(date(2032,4,3),[r['day'] for r in rows])
    def test_identical_duplicate_is_deduplicated(self):
        _,rows,notes=self.engine.stage_run(dict(self.payload,scenario='duplicate'));self.assertEqual(len(rows),7);self.assertIn('duplicado',notes[0])
    def test_conflict_stops_generation(self):
        with self.assertRaises(InputFault):self.engine.produce_run(dict(self.payload,scenario='conflict'))
        self.assertEqual(self.engine.timeline(),[]);self.assertEqual(list(self.db.outputs.iterdir()),[])
    def test_warning_keeps_numeric_measurements(self):
        _,rows,notes=self.engine.stage_run(dict(self.payload,scenario='warning'));self.assertEqual(len(rows),7);self.assertTrue(notes)
    def test_source_rejects_unknown_scenario(self):
        with self.assertRaises(InputFault):self.engine.preview_run(dict(self.payload,scenario='remote'))
    def test_sqlite_created_automatically_in_new_directory(self):
        destination=self.folder/'new-installation'
        self.assertFalse(destination.exists())
        fresh=LabDatabase(destination)
        self.assertTrue(fresh.path.is_file())
        self.assertEqual(len(fresh.catalog()['units']),6)
    def test_generation_works_with_all_network_calls_forbidden(self):
        with patch('socket.socket.connect',side_effect=AssertionError('Network forbidden')), patch('socket.getaddrinfo',side_effect=AssertionError('DNS forbidden')):
            job=self.engine.produce_run(self.payload)
            book=load_workbook(BytesIO(self.engine.obtain_file(job['id'])))
            try:self.assertEqual(book.sheetnames,['Reporte','Resumen'])
            finally:book.close()
    def test_application_modules_belong_to_this_project(self):
        import sys
        root=Path(__file__).resolve().parents[1]
        for name,module in tuple(sys.modules.items()):
            if name=='fleet_lab' or name.startswith('fleet_lab.'):
                self.assertTrue(Path(module.__file__).resolve().is_relative_to(root),name)
    def test_foreign_measurements_rejected(self):
        units=[self.db.eligible('boreal','2032-04-01','2032-04-01')[0]];day=date(2032,4,1)
        with self.assertRaises(InputFault):normalize_observations([dict(unit_id='foreign',day=day,distance=1,energy=1)],units,[day])

class CatalogRecovery(IsolatedLab):
    def test_inactive_and_future_association_are_not_eligible(self):
        self.assertEqual([u['id'] for u in self.db.eligible('prism','2032-04-01','2032-04-07')],['veh-903'])
        self.assertEqual([u['id'] for u in self.db.eligible('orbit','2032-04-01','2032-04-07')],['veh-906'])
    def test_vehicle_is_not_automatically_assigned(self):
        unit=dict(plate='DEMO-980',internal_key='LAB-7980',lookup_key='lookup-lab-980',label='Unidad nueva',active=True)
        identity=self.db.save_unit(unit);self.assertNotIn(identity,[u['id'] for u in self.db.eligible('boreal','2032-04-01','2032-04-07')])
    def test_explicit_assignment_controls_selection(self):
        self.db.link_unit(dict(unit_id='veh-903',series_id='boreal',valid_from='2032-04-01',valid_to='2032-04-07'))
        self.assertIn('veh-903',[u['id'] for u in self.db.eligible('boreal','2032-04-01','2032-04-07')])
    def test_partial_validity_does_not_cover_whole_period(self):
        self.db.link_unit(dict(unit_id='veh-901',series_id='boreal',valid_from='2032-04-03',valid_to='2032-04-07'))
        with self.assertRaises(InputFault):self.engine.produce_run(self.payload)
    def test_duplicate_plate_is_rejected(self):
        unit=dict(self.db.catalog()['units'][0]);unit.pop('id')
        with self.assertRaises(InputFault):self.db.save_unit(unit)
    def test_preview_import_is_read_only(self):
        before=self.db.catalog();data=deepcopy(before);data['units'][0]['label']='Nombre alternativo'
        self.db.validate_catalog(data);self.assertEqual(self.db.catalog(),before)
    def test_import_rejects_extra_fields_before_backup(self):
        data=self.db.catalog();data['private']='not part of the schema'
        with self.assertRaises(InputFault):self.db.import_catalog(data)
        self.assertEqual(list(self.db.archives.iterdir()),[])
    def test_backup_restore_retains_reports_and_audit(self):
        job=self.engine.produce_run(self.payload);before=self.db.catalog();backup=self.db.archive_catalog()
        unit=dict(before['units'][0],label='Unidad modificada');self.db.save_unit(unit)
        self.db.restore_catalog(backup['name'],backup['sha256'])
        self.assertEqual(self.db.catalog(),before);self.assertEqual(len(self.engine.timeline()),1)
        self.assertTrue(self.engine.obtain_file(job['id']));self.assertIn('unit_saved',[e['kind'] for e in self.db.events()])
        self.assertIn('catalog_restored',[e['kind'] for e in self.db.events()])
    def test_tampered_backup_is_rejected(self):
        backup=self.db.archive_catalog();(self.db.archives/backup['name']).write_text('{}')
        with self.assertRaises(InputFault):self.db.restore_catalog(backup['name'],backup['sha256'])
    def test_backup_path_traversal_is_rejected(self):
        with self.assertRaises(InputFault):self.db.restore_catalog('../seed.json','bad')
    def test_restart_retains_catalog_without_duplicates(self):
        restarted=LabDatabase(self.folder);self.assertEqual(restarted.catalog(),self.db.catalog());self.assertEqual(len(restarted.catalog()['units']),6)
    def test_future_schema_not_silently_downgraded(self):
        with self.db.transaction() as db:db.execute('PRAGMA user_version=5')
        with self.assertRaises(InputFault):LabDatabase(self.folder)
    def test_import_validates_foreign_keys(self):
        data=self.db.catalog();data['assignments'][0]['unit_id']='absent'
        with self.assertRaises(InputFault):self.db.validate_catalog(data)

class ExportAndDecisions(IsolatedLab):
    def workbook(self,payload):
        job=self.engine.produce_run(payload);book=load_workbook(BytesIO(self.engine.obtain_file(job['id'])))
        self.addCleanup(book.close);return job,book
    def test_regular_excel_has_only_two_sheets(self):
        _,book=self.workbook(self.payload);self.assertEqual(book.sheetnames,['Reporte','Resumen'])
    def test_observations_add_warnings_sheet(self):
        _,book=self.workbook(dict(self.payload,scenario='mixed'));self.assertEqual(book.sheetnames,['Reporte','Resumen','Advertencias'])
    def test_confirmed_zero_does_not_add_unnecessary_warning_sheet(self):
        _,book=self.workbook(dict(self.payload,scenario='zero'))
        self.assertEqual(book.sheetnames,['Reporte','Resumen'])
    def test_advanced_output_is_explicit_opt_in(self):
        _,book=self.workbook(dict(self.payload,advanced=True));self.assertEqual(book.sheetnames,['Reporte','Resumen','Detalle técnico'])
    def test_dates_numbers_tables_and_metadata(self):
        _,book=self.workbook(self.payload);sheet=book['Reporte'];self.assertEqual(sheet['C2'].data_type,'d');self.assertEqual(sheet['D2'].data_type,'n');self.assertEqual(len(sheet.tables),1)
        self.assertEqual(book.properties.creator,'Fleet Energy Lab');self.assertEqual(sheet.freeze_panes,'A2')
    def test_formula_like_identifier_written_as_text(self):
        unit=next(u for u in self.db.catalog()['units'] if u['id']=='veh-901');self.db.save_unit(dict(unit,internal_key='=2+3'))
        _,book=self.workbook(self.payload);self.assertEqual(book['Reporte']['B2'].data_type,'s');self.assertTrue(book['Reporte']['B2'].value.startswith("'="))
    def test_preview_does_not_create_report(self):
        view=self.engine.preview_run(self.payload);self.assertEqual(view['count'],7);self.assertEqual(self.engine.timeline(),[])
    def test_history_uses_immutable_catalog_snapshot(self):
        job,_=self.workbook(self.payload);unit=next(u for u in self.db.catalog()['units'] if u['id']=='veh-901');self.db.save_unit(dict(unit,plate='DEMO-999'))
        self.assertEqual(self.engine.timeline()[0]['units'],['DEMO-901']);self.assertTrue(self.engine.obtain_file(job['id']))
    def test_exact_hash_required_for_decision(self):
        job,_=self.workbook(self.payload)
        with self.assertRaises(InputFault):self.engine.record_decision(dict(id=job['id'],sha256='wrong',decision='accepted'))
        self.engine.record_decision(dict(id=job['id'],sha256=job['sha256'],decision='accepted'))
        self.assertEqual(self.engine.timeline()[0]['decision'],'accepted');self.assertEqual(self.db.events()[0]['detail']['sha256'],job['sha256'])
    def test_changed_report_cannot_be_downloaded_or_accepted(self):
        job,_=self.workbook(self.payload);(self.db.outputs/(job['id']+'.xlsx')).write_bytes(b'changed')
        with self.assertRaises(InputFault):self.engine.obtain_file(job['id'])
        with self.assertRaises(InputFault):self.engine.record_decision(dict(id=job['id'],sha256=job['sha256'],decision='accepted'))
    def test_repeated_units_and_inactive_unit_rejected(self):
        for units in (['veh-901','veh-901'],['veh-904'],[]):
            with self.subTest(units=units),self.assertRaises(InputFault):self.engine.produce_run(dict(self.payload,unit_ids=units))

class HttpBoundary(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='fleet-http-');self.server=create_lab_server(self.temp.name)
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start();self.port=self.server.server_port
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.thread.join(3);self.temp.cleanup()
    def call(self,method,path,data=None,headers=None):
        h={'Host':f'127.0.0.1:{self.port}'}
        if data is not None:h.update({'Content-Type':'application/json','Origin':f'http://127.0.0.1:{self.port}'})
        h.update(headers or {});conn=http.client.HTTPConnection('127.0.0.1',self.port,timeout=5)
        try:
            conn.request(method,path,json.dumps(data) if data is not None else None,h);response=conn.getresponse();return response.status,dict(response.getheaders()),response.read()
        finally:conn.close()
    def test_listener_is_loopback_only(self):self.assertEqual(self.server.server_address[0],'127.0.0.1')
    def test_home_and_security_headers(self):
        status,headers,body=self.call('GET','/');self.assertEqual(status,200);self.assertIn("default-src 'self'",headers['Content-Security-Policy']);self.assertIn(b'Fleet Energy',body)
    def test_http_runs_when_only_loopback_connections_are_permitted(self):
        original_connect=socket.socket.connect
        def loopback_only(connection,address):
            if not isinstance(address,tuple) or address[0]!='127.0.0.1':
                raise AssertionError('Only loopback is permitted')
            return original_connect(connection,address)
        with patch('socket.socket.connect',loopback_only):
            self.assertEqual(self.call('GET','/api/state')[0],200)
    def test_unknown_host_rejected(self):self.assertEqual(self.call('GET','/api/state',headers={'Host':'127.0.0.2:1'})[0],403)
    def test_foreign_and_missing_origin_rejected(self):
        self.assertEqual(self.call('POST','/api/backup',{},headers={'Origin':'http://127.0.0.2:1'})[0],403)
        self.assertEqual(self.call('POST','/api/backup',{},headers={'Origin':''})[0],403)
    def test_wrong_content_type_rejected(self):self.assertEqual(self.call('POST','/api/backup',{},headers={'Content-Type':'text/plain'})[0],415)
    def test_static_traversal_unavailable(self):self.assertEqual(self.call('GET','/../datos_demo/seed.json')[0],404)
    def test_state_contains_no_absolute_storage_path(self):
        status,_,body=self.call('GET','/api/state');self.assertEqual(status,200);self.assertNotIn(self.temp.name.encode(),body);self.assertNotIn(b'Set-Cookie',body)
    def test_generate_then_download(self):
        payload=dict(series_id='boreal',unit_ids=['veh-901'],start='2032-04-01',end='2032-04-02',scenario='steady',advanced=False)
        status,_,body=self.call('POST','/api/generate',payload);self.assertEqual(status,200);identity=json.loads(body)['id']
        status,headers,content=self.call('GET','/download/'+identity);self.assertEqual(status,200);self.assertIn('attachment',headers['Content-Disposition']);self.assertTrue(content.startswith(b'PK'))
    def test_backup_requires_empty_body(self):self.assertEqual(self.call('POST','/api/backup',{'unexpected':1})[0],422)
    def test_unregistered_download_is_rejected(self):self.assertEqual(self.call('GET','/download/'+'0'*32)[0],422)

class WindowLifecycle(unittest.TestCase):
    def exercise(self, connected, closed):
        from playwright.sync_api import Error
        engine=MagicMock();browser=engine.chromium.launch.return_value
        browser.is_connected.return_value=connected
        page=browser.new_context.return_value.new_page.return_value
        page.is_closed.return_value=closed
        page.goto.side_effect=Error('Synthetic interrupted navigation')
        manager=MagicMock();manager.__enter__.return_value=engine
        return manager,browser
    def test_close_during_initial_navigation_is_normal(self):
        from fleet_lab.web import open_ephemeral_window
        manager,browser=self.exercise(True,True)
        with patch('playwright.sync_api.sync_playwright',return_value=manager):open_ephemeral_window('http://127.0.0.1:1')
        browser.close.assert_called_once()
    def test_disconnected_browser_is_already_closed(self):
        from fleet_lab.web import open_ephemeral_window
        manager,browser=self.exercise(False,True)
        with patch('playwright.sync_api.sync_playwright',return_value=manager):open_ephemeral_window('http://127.0.0.1:1')
        browser.close.assert_not_called()
    def test_navigation_error_with_open_window_remains_an_error(self):
        from playwright.sync_api import Error
        from fleet_lab.web import open_ephemeral_window
        manager,browser=self.exercise(True,False)
        with patch('playwright.sync_api.sync_playwright',return_value=manager),self.assertRaises(Error):open_ephemeral_window('http://127.0.0.1:1')
        browser.close.assert_called_once()

if __name__=='__main__':unittest.main()
